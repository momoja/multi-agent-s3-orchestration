import os
import json
import time

from strands import Agent
from strands.models.bedrock import BedrockModel
from docx import Document

MOUNT = "/mnt/s3files"
MODEL_ID = os.environ.get("MODEL_ID", "us.anthropic.claude-sonnet-4-6")
REGION = os.environ.get("AWS_DEFAULT_REGION", "us-east-1")
PROCESSED_PATH = f"{MOUNT}/.processed-ecs.json"

model = BedrockModel(model_id=MODEL_ID, region_name=REGION)
agent = Agent(
    model=model,
    system_prompt="""
You are a report assembly agent. Read validated compliance documents, aggregate
findings, and produce an executive report. Return JSON with: batch_summary,
high_risk_items, action_items, executive_narrative.
""",
)


def load_processed() -> set:
    if os.path.exists(PROCESSED_PATH):
        with open(PROCESSED_PATH) as f:
            return set(json.load(f))
    return set()


def save_processed(processed: set) -> None:
    with open(PROCESSED_PATH, "w") as f:
        json.dump(list(processed), f)


def main() -> None:
    processed = load_processed()
    print(f"report-agent starting, mount={MOUNT}", flush=True)

    while True:
        validated_dir = f"{MOUNT}/validated"
        if not os.path.isdir(validated_dir):
            time.sleep(15)
            continue

        files = [f for f in os.listdir(validated_dir) if f.endswith(".json") and f not in processed]
        for validated_file in files:
            try:
                handle_file(validated_dir, validated_file)
            except Exception as exc:  # noqa: BLE001
                print(f"error processing {validated_file}: {exc}", flush=True)
                continue
            processed.add(validated_file)
            save_processed(processed)

        time.sleep(15)


def handle_file(validated_dir: str, validated_file: str) -> None:
    with open(f"{validated_dir}/{validated_file}") as f:
        doc = json.load(f)

    result = agent(f"Generate executive compliance report: {json.dumps(doc['content'])}")
    envelope = {
        "id": doc["id"],
        "stage": "report",
        "type": doc["type"],
        "content": json.loads(result.message["content"][0]["text"]),
    }
    with open(f"{MOUNT}/reports/report-{doc['id']}.json", "w") as f:
        json.dump(envelope, f, indent=2)

    report_doc = Document()
    report_doc.add_heading("Compliance Report", 0)
    report_doc.add_paragraph(envelope["content"]["executive_narrative"])
    report_doc.save(f"{MOUNT}/documents/report-{doc['id']}.docx")
    print(f"wrote report-{doc['id']}", flush=True)


if __name__ == "__main__":
    main()
